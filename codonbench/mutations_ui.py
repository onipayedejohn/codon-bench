"""The Mutations page: one variant, two sequences, a whole file of variants, and structural changes.

Works on any reference: a coding sequence (c. notation), an unannotated DNA sequence (g.), or an
annotated GenBank record of any organism (g., or m. for mitochondria), where every variant is placed in
its gene and translated with that gene's genetic code.
"""
from __future__ import annotations

import asyncio
import csv
import tempfile
from pathlib import Path

import pandas as pd
from shiny import reactive, render, req, ui

from . import demos, ml
from . import files as F
from . import genome as G
from . import mutations as M
from . import seq as S
from . import viz as V
from .mutation_views import (_badge, _chromosome, _dna_compare, _effect_view, _facts, _fusion_chips, _notice,
                             _wrap)

_plot = V.as_plot
SAMPLES = Path(__file__).resolve().parent.parent / "data" / "samples"

REF_CHOICES = {"hbb": "Human beta-globin, HBB (coding sequence)",
               "adamts1": "Human ADAMTS1 gene with introns (GRCh38)",
               "mito": "Human mitochondrial genome (rCRS)",
               "sars2": "SARS-CoV-2 genome (Wuhan-Hu-1)",
               "own": "My own reference"}
REF_FILES = {"adamts1": "references/human_ADAMTS1_genomic.gb", "mito": "references/human_mitochondrion_rCRS.gb",
             "sars2": "references/sars_cov_2_reference.gb"}
EXAMPLES = {
    "hbb": {k: f"{k}  {v.split(':')[0]}" for k, v in M.HBB_EXAMPLES.items()},
    "adamts1": {"g.686A>G": "g.686A>G  missense in exon 1", "g.692A>T": "g.692A>T  nonsense in exon 1",
                "g.1386G>A": "g.1386G>A  splice donor +1, intron 1", "g.1390A>G": "g.1390A>G  splice region +5",
                "g.2923A>G": "g.2923A>G  splice acceptor -2, intron 1", "g.2000G>A": "g.2000G>A  deep intronic",
                "g.3001del": "g.3001del  frameshift in exon 2"},
    "mito": {"m.3243A>G": "m.3243A>G  MT-TL1, MELAS", "m.11778G>A": "m.11778G>A  MT-ND4 R340H, LHON",
             "m.14484T>C": "m.14484T>C  MT-ND6 M64V, LHON (minus-strand gene)",
             "m.8993T>G": "m.8993T>G  MT-ATP6 L156R, NARP", "m.1555A>G": "m.1555A>G  MT-RNR1, aminoglycoside deafness",
             "m.7028C>T": "m.7028C>T  MT-CO1 synonymous, non-H haplogroup marker"},
    "sars2": {"A23403G": "A23403G  S:D614G", "A23063T": "A23063T  S:N501Y", "G23012A": "G23012A  S:E484K",
              "C23604A": "C23604A  S:P681H", "g.21765_21770del": "g.21765_21770del  S:H69-V70 deletion",
              "C27972T": "C27972T  ORF8:Q27 stop", "C14408T": "C14408T  ORF1b:P314L (RdRp P323L)"},
    "own": {"": "Type a change below"},
}
BATCH_DEMOS = {  # button id -> (label, sample file, reference)
    "md_hbb_csv": ("HBB list (CSV)", "variants/hbb_variants_hgvs.csv", "hbb"),
    "md_hbb_xlsx": ("HBB list (Excel)", "variants/hbb_variants_hgvs.xlsx", "hbb"),
    "md_hbb_fa": ("HBB mutants (FASTA)", "variants/hbb_mutant_sequences.fasta", "hbb"),
    "md_sars_vcf": ("SARS-CoV-2 Alpha (VCF)", "variants/sars_cov_2_alpha_variants.vcf", "sars2"),
    "md_sars_fa": ("SARS-CoV-2 consensus genome (FASTA)", "variants/sars_cov_2_alpha_like_consensus.fasta", "sars2"),
    "md_mito_vcf": ("Mitochondrial variants (VCF)", "variants/human_mtdna_variants.vcf", "mito"),
    "md_ad_csv": ("ADAMTS1 splice variants (CSV)", "variants/adamts1_genomic_variants.csv", "adamts1"),
}
PARTNER = ("ATGGATACCGAGCGTAAACTGCAGTCTGGCAAACGTCCGGTTGAAGATCTGCGCTTCAAAGGTGCGGAAATCCTGTCTGAACGTGGCATGCAG"
           "CTGAAATAA")   # a synthetic partner gene for the translocation example
SPLICE_BATCH_LIMIT = 500


def builtin_reference(key: str) -> G.Reference:
    if key == "hbb":
        return G.Reference("HBB", "Human beta-globin coding sequence (NM_000518.5)", demos.HBB_CDS, [], "coding",
                           "c", "Homo sapiens", 1, True)
    return G.load((SAMPLES / REF_FILES[key]).read_text())[0]


# ---------------------------------------------------------------------------------------------
# UI
# ---------------------------------------------------------------------------------------------
def mutation_ui():
    return ui.nav_panel(
        "Mutations",
        ui.layout_sidebar(
            ui.sidebar(
                ui.h3("Reference"),
                ui.input_select("m_ref", "Compare variants against", REF_CHOICES, selected="hbb"),
                ui.panel_conditional(
                    "input.m_ref === 'own'",
                    ui.input_text_area("m_ref_text", "Paste FASTA, GenBank or plain DNA", rows=6, width="100%",
                                       placeholder=">my_gene\nATGGCC...  or a GenBank record starting with LOCUS"),
                    ui.input_file("m_ref_file", "Or upload a reference", accept=[".fasta", ".fa", ".fna", ".txt",
                                                                                ".gb", ".gbk", ".gbff", ".gz"],
                                  button_label="Choose file", placeholder="FASTA or GenBank"),
                    ui.input_action_button("m_ref_use", "Use this reference", class_="btn-ghost w-100 mb-3"),
                ),
                ui.input_select("m_code", "Genetic code (when the reference does not say)",
                                {str(k): v for k, v in S.GENETIC_CODES.items()}, selected="1"),
                ui.input_switch("m_splice", "Predict splice impact (ML)", True),
                ui.output_ui("m_ref_check"),
                width=340, open={"desktop": "open", "mobile": "always"},
            ),
            ui.navset_card_underline(
                ui.nav_panel("Single variant", _single_ui()),
                ui.nav_panel("Compare sequences", _compare_ui()),
                ui.nav_panel("Batch variants", _batch_ui()),
                ui.nav_panel("Structural changes", _structural_ui()),
                id="m_tabs",
            ),
        ),
    )


def _single_ui():
    return ui.TagList(
        ui.layout_columns(
            ui.div(ui.input_select("m_example", "Known variants for this reference", EXAMPLES["hbb"],
                                   selected="c.20A>T"),
                   ui.output_ui("m_example_note")),
            ui.div(
                ui.layout_columns(
                    ui.input_select("m_type", "Build a change", {"substitution": "Substitution",
                                                                 "insertion": "Insertion", "deletion": "Deletion",
                                                                 "duplication": "Duplication",
                                                                 "inversion": "Inversion"}, selected="substitution"),
                    ui.input_numeric("m_pos", "Position", 20, min=1, step=1),
                    ui.div(
                        ui.panel_conditional("input.m_type === 'substitution'",
                                             ui.input_select("m_base", "New base", {b: b for b in "ACGT"},
                                                             selected="T")),
                        ui.panel_conditional("input.m_type === 'insertion'",
                                             ui.input_text("m_ins", "Bases to insert", value="T")),
                        ui.panel_conditional("['deletion','duplication','inversion'].includes(input.m_type)",
                                             ui.input_numeric("m_len", "Length (bases)", 3, min=1, step=1))),
                    col_widths=(4, 4, 4)),
                ui.input_text("m_hgvs", "Or type a change", value="c.20A>T",
                              placeholder="c.20A>T, g.1386G>A, m.3243A>G, A23403G, g.100_105del"),
                ui.input_action_button("m_apply", "Analyse variant", class_="btn-primary"),
            ),
            col_widths=(5, 7)),
        ui.hr(),
        ui.output_ui("m_result"),
    )


def _compare_ui():
    return ui.TagList(
        ui.p("Paste a mutant gene or a whole genome (for example a consensus sequence). Every difference from the "
             "reference is found, named in HGVS notation and placed in its gene.", class_="muted"),
        ui.input_text_area("m_mutant", "Mutant sequence or genome (FASTA or plain letters)", rows=6, width="100%",
                           value=_wrap(M.apply(demos.HBB_CDS, M.parse_hgvs("c.20A>T", demos.HBB_CDS)))),
        ui.div(ui.input_action_button("m_compare", "Compare with the reference", class_="btn-primary"),
               ui.input_action_button("m_cmp_demo", "HBB frameshift example", class_="btn-ghost ms-2"),
               ui.input_action_button("m_cmp_demo2", "SARS-CoV-2 consensus genome", class_="btn-ghost ms-2"),
               class_="mb-3 demo-row"),
        ui.output_ui("m_compare_out"),
    )


def _batch_ui():
    return ui.TagList(
        ui.p("Upload variants in a VCF, a CSV or Excel table (a column of HGVS changes, or pos, ref and alt columns), or "
             "a FASTA file of mutant sequences or genomes. Each variant is checked against the reference, annotated and "
             f"counted. Files up to {F.MAX_BYTES // 1048576} MB.", class_="muted"),
        ui.layout_columns(
            ui.input_file("m_file", "Upload variants", accept=[".vcf", ".csv", ".tsv", ".xlsx", ".fasta", ".fa", ".fna",
                                                               ".txt", ".gz"],
                          button_label="Choose file", placeholder="VCF, CSV, XLSX, FASTA"),
            ui.div(ui.p("Or try an example (the reference switches to match)", class_="form-label mb-1"),
                   ui.div(*[ui.input_action_button(k, v[0], class_="btn-ghost") for k, v in BATCH_DEMOS.items()],
                          class_="demo-row")),
            col_widths=(4, 8)),
        ui.output_ui("m_batch_out"),
    )


def _structural_ui():
    return ui.TagList(
        ui.p("Large rearrangements move whole segments. They are drawn to scale, then checked base by base for their "
             "effect on the gene. For an annotated genome, the longest coding sequence is used.", class_="muted"),
        ui.navset_pill(
            ui.nav_panel(
                "Inversion",
                ui.layout_columns(ui.input_numeric("inv_a", "From (c.)", 31, min=1),
                                  ui.input_numeric("inv_b", "To (c.)", 60, min=2), col_widths=(3, 3)),
                ui.output_ui("inv_out")),
            ui.nav_panel(
                "Translocation",
                ui.input_text_area("tr_b", "Partner gene B (coding sequence)", value=_wrap(PARTNER), rows=3,
                                   width="100%"),
                ui.layout_columns(ui.input_numeric("tr_a", "Break after base (gene A)", 30, min=1),
                                  ui.input_numeric("tr_bb", "Break after base (gene B)", 9, min=1),
                                  col_widths=(4, 4)),
                ui.output_ui("tr_out")),
            ui.nav_panel(
                "Duplication and deletion",
                ui.layout_columns(ui.input_numeric("dd_a", "From (c.)", 31, min=1),
                                  ui.input_numeric("dd_b", "To (c.)", 60, min=1),
                                  ui.input_radio_buttons("dd_kind", "Change", {"duplication": "Duplicate",
                                                                               "deletion": "Delete"}, inline=True),
                                  col_widths=(3, 3, 6)),
                ui.output_ui("dd_out")),
            id="struct_tabs"),
    )


# ---------------------------------------------------------------------------------------------
# Server
# ---------------------------------------------------------------------------------------------
def mutation_server(input, output, session, send_to_analyzer, log=lambda *a, **k: None):
    own = reactive.value(None)          # (Reference | None, problem)
    chosen = reactive.value("c.20A>T")
    quiet = {"until": 0.0}              # ignore builder events caused by the server syncing the builder

    def sync_builder(text: str):
        r = ref()
        if r is None or not text:
            return
        try:
            v = G.parse_genomic_change(text, r)
        except M.BadNotation:
            return
        import time
        quiet["until"] = time.time() + 1.5
        p0, rr, aa = v.core()
        ui.update_numeric("m_pos", value=p0 + 1 if rr else p0)
        if len(rr) == 1 and len(aa) == 1:
            ui.update_select("m_type", selected="substitution")
            ui.update_select("m_base", selected=aa)

    def _load_own(text: str, name: str):
        try:
            refs = G.load(text, name)
            note = f" The file has {len(refs)} records; the first is used." if len(refs) > 1 else ""
            own.set((refs[0], note))
        except M.BadNotation as e:
            own.set((None, str(e)))

    @reactive.effect
    @reactive.event(input.m_ref_file)
    def _own_file():
        f = input.m_ref_file()
        if not f:
            return
        try:
            kind = F.sniff(f[0]["datapath"], f[0]["name"], f[0]["size"])
            if kind.family != "sequence":
                raise F.Blocked("A reference must be a FASTA or GenBank file.")
            F.check_text_file(f[0]["datapath"], kind)
            _load_own(G.read_file(f[0]["datapath"], f[0]["name"]), f[0]["name"])
        except F.Blocked as e:
            own.set((None, str(e)))

    @reactive.effect
    @reactive.event(input.m_ref_use)
    def _own_text():
        if (input.m_ref_text() or "").strip():
            _load_own(input.m_ref_text(), "pasted reference")

    @reactive.calc
    def ref_pair() -> tuple[G.Reference | None, str]:
        key = input.m_ref()
        if key != "own":
            r = builtin_reference(key)
        else:
            r, note = own() or (None, "Paste or upload a reference, then press 'Use this reference'.")
            if r is None:
                return None, note
        if r.kind != "annotated":
            r.table = int(input.m_code())
        return r, ""

    def ref() -> G.Reference | None:
        return ref_pair()[0]

    @reactive.calc
    def splice_on() -> bool:
        r = ref()
        return bool(input.m_splice()) and r is not None and r.eukaryotic_nuclear and ml.available()

    @render.ui
    def m_ref_check():
        r, prob = ref_pair()
        if r is None:
            return _notice("blocked" if "not" in prob.lower() else "empty", "No reference yet", prob)
        kind = {"coding": "Coding sequence (c. numbering)", "genomic": "DNA sequence without annotation (g.)",
                "annotated": "Annotated genome (" + ("m." if r.prefix == "m" else "g.") + " numbering)"}[r.kind]
        items = [("Reference", r.id), ("Type", kind), ("Size", r.summary())]
        if r.organism:
            items.append(("Organism", r.organism))
        items.append(("Genetic code", S.GENETIC_CODES.get(r.table, str(r.table))))
        if not r.eukaryotic_nuclear:
            items.append(("Splice prediction", "Off: no spliceosomal introns in this genome"))
        elif not ml.available():
            items.append(("Splice prediction", "Models not installed"))
        return ui.div(_facts(items), class_="ref-card")

    # ---- examples -----------------------------------------------------------------------------------
    @reactive.effect
    @reactive.event(input.m_ref)
    def _ref_changed():
        ex = EXAMPLES.get(input.m_ref(), EXAMPLES["own"])
        first = next(iter(ex))
        ui.update_select("m_example", choices=ex, selected=first)
        if first:
            ui.update_text("m_hgvs", value=first)
            chosen.set(first)
            sync_builder(first)

    @reactive.effect
    @reactive.event(input.m_example, ignore_init=True)
    def _ex():
        if input.m_example():
            ui.update_text("m_hgvs", value=input.m_example())
            chosen.set(input.m_example())
            sync_builder(input.m_example())

    @render.ui
    def m_example_note():
        if input.m_ref() == "hbb":
            return ui.p(M.HBB_EXAMPLES.get(input.m_example(), ""), class_="small-note mt-1")
        r = ref()
        if r is None:
            return None
        return ui.p(f"Positions are numbered from 1 along {r.id}. You can also type changes such as "
                    f"{r.prefix}.100A>G, {r.prefix}.100_105del, {r.prefix}.100_101insT or {r.prefix}.100_102dup.",
                    class_="small-note mt-1")

    @reactive.effect
    @reactive.event(input.m_type, input.m_pos, input.m_base, input.m_ins, input.m_len, ignore_init=True)
    def _build():
        import time
        if time.time() < quiet["until"]:
            return
        r = ref()
        if r is None or not input.m_pos():
            return
        dna, pre = r.seq, r.prefix
        p, t = int(input.m_pos()), input.m_type()
        if not 1 <= p <= len(dna):
            return
        if t == "substitution":
            if input.m_base() == dna[p - 1]:
                ui.update_select("m_base", selected=next(b for b in "TACG" if b != dna[p - 1]))
                return
            txt = f"{pre}.{p}{dna[p - 1]}>{input.m_base()}"
        elif t == "insertion":
            p = min(p, len(dna) - 1)
            ins = "".join(c for c in (input.m_ins() or "").upper() if c in "ACGT") or "T"
            txt = f"{pre}.{p}_{p + 1}ins{ins}"
        else:
            n = max(1, int(input.m_len() or 1))
            end = min(len(dna), p + n - 1)
            if t == "inversion" and end == p:
                end = min(len(dna), p + 1)
            suffix = {"deletion": "del", "duplication": "dup", "inversion": "inv"}[t]
            txt = f"{pre}.{p}" + (f"_{end}" if end != p else "") + suffix
        ui.update_text("m_hgvs", value=txt)
        chosen.set(txt)

    @reactive.effect
    @reactive.event(input.m_apply)
    def _apply():
        chosen.set(input.m_hgvs())

    # ---- single variant -------------------------------------------------------------------------
    @render.ui
    def m_result():
        r, prob = ref_pair()
        if r is None:
            return _notice("empty", "Choose a reference first", prob)
        text = (chosen() or "").strip()
        if not text:
            return _notice("empty", "Type a change", "For example g.100A>G, or build one with the controls above.")
        try:
            v = G.parse_genomic_change(text, r)
        except M.BadNotation as e:
            return _notice("blocked", "That change could not be read", str(e))
        a = G.annotate(r, v, splice=splice_on())
        if not a["status"].startswith("Annotated"):
            return _notice("blocked", "That change does not fit the reference", a["status"])
        log("Mutation", f"{a['hgvs_g'] or text} on {r.id}", f"{a['effect']} {a['hgvs_p']}".strip())
        parts = []
        typed = text.replace(" ", "")
        canonical = a["hgvs_c"] if r.kind == "coding" else a["hgvs_g"]
        if typed.startswith(("c.", "g.", "m.")) and canonical and typed != canonical:
            parts.append(_notice("warn", "Rewritten in HGVS form",
                                 f"{typed} is written {canonical} in HGVS: insertions and deletions are placed as far "
                                 "3' as possible, and a copied stretch is a duplication. Both describe the same "
                                 "sequence."))
        parts.append(_summary_card(r, a))
        detail = G.cds_effect(r, v)
        if detail:
            f, cds, mut, m, e = detail
            parts.append(_effect_view(cds, mut, m, e, is_hbb=(r.id == "HBB")))
        else:
            parts.append(_noncoding_card(r, v, a))
        if splice_on() and a.get("splice_detail"):
            parts.append(_splice_card(a))
            parts.append(ui.output_ui("m_splice_plot"))
        elif input.m_splice() and r is not None and not r.eukaryotic_nuclear:
            parts.append(ui.p("Splice prediction is off for this reference: viral, bacterial and organelle genomes "
                              "do not have spliceosomal introns.", class_="small-note"))
        return ui.TagList(*parts)

    @render.ui
    @_plot
    def m_splice_plot():
        r = ref()
        req(r is not None and splice_on())
        try:
            v = G.parse_genomic_change((chosen() or "").strip(), r)
        except M.BadNotation:
            req(False)
        a = G.annotate(r, v, splice=True)
        req(a.get("splice_detail"))
        return V.splice_profile_plot(a["splice_detail"]["profile"], v.pos, r.prefix)

    # ---- compare ----------------------------------------------------------------------------------
    cmp_job = reactive.value(None)

    @reactive.effect
    @reactive.event(input.m_cmp_demo)
    def _cmp_demo():
        mutant = _wrap(M.apply(demos.HBB_CDS, M.parse_hgvs("c.126_129del", demos.HBB_CDS)))
        ui.update_select("m_ref", selected="hbb")
        ui.update_text_area("m_mutant", value=mutant)
        cmp_job.set(("hbb", mutant))

    @reactive.effect
    @reactive.event(input.m_cmp_demo2)
    def _cmp_demo2():
        text = (SAMPLES / "variants" / "sars_cov_2_alpha_like_consensus.fasta").read_text()
        ui.update_select("m_ref", selected="sars2")
        ui.update_text_area("m_mutant", value=text[:3000] + ("\n..." if len(text) > 3000 else ""))
        cmp_job.set(("sars2", text))

    @reactive.effect
    @reactive.event(input.m_compare)
    def _cmp():
        cmp_job.set((None, input.m_mutant()))

    @render.ui
    def m_compare_out():
        job = cmp_job()
        if job is None:
            return _notice("empty", "Paste a sequence and press Compare",
                           "Or load one of the examples. A gene is aligned base by base; a genome is compared by "
                           "anchoring exact matches first, so a 30 kb virus takes under a second.")
        key, text = job
        with reactive.isolate():
            r = builtin_reference(key) if key else ref()
        if r is None:
            return _notice("blocked", "Choose a reference first", "")
        body = text.split("\n", 1)[1] if text.lstrip().startswith(">") else text
        p = S.parse(body)
        if p.kind not in ("DNA", "RNA"):
            return _notice("blocked", "The sequence was not analysed", p.reason or "It must be DNA or RNA.")
        q = S.to_dna(p.seq)
        try:
            if len(r.seq) <= 20000 and len(q) <= 20000 and r.kind != "annotated":
                variants = [G.mutation_to_variant(m, r) for m in M.detect(r.seq, q)]
            else:
                variants = G.compare_genomes(r, q)
        except M.BadNotation as e:
            return _notice("blocked", "Could not compare", str(e))
        if not variants:
            return _notice("ok", "No differences", "The sequence matches the reference.")
        rows = [G.annotate(r, v) for v in variants[:500]]
        log("Compare", f"{len(variants)} differences vs {r.id}", ", ".join(x["hgvs_g"] for x in rows[:4]))
        df = pd.DataFrame([{"Change": x["hgvs_g"], "Type": x["type"], "Gene": x["gene"], "c.": x["hgvs_c"],
                            "Protein": x["protein_short"] or x["hgvs_p"], "Effect": x["effect"]} for x in rows])
        table = ui.div(ui.HTML(df.to_html(index=False, classes="table table-sm diff-table", border=0, justify="left",
                                          escape=True)), class_="table-scroll")
        out = [ui.p(f"{len(variants)} difference{'s' if len(variants) != 1 else ''} found.", class_="muted"), table]
        if r.kind == "coding" and len(variants) >= 1:
            e = M.effect(r.seq, q, M.detect(r.seq, q)[0] if len(variants) == 1 else None, r.table)
            m1 = M.detect(r.seq, q)[0] if len(variants) == 1 else None
            out.append(ui.h5("Combined effect on the protein", class_="mt-3"))
            out.append(_effect_view(r.seq, q, m1, e, is_hbb=(r.id == "HBB")))
        return ui.TagList(*out)

    # ---- batch ------------------------------------------------------------------------------------
    job = reactive.value(None)

    @reactive.effect
    @reactive.event(input.m_file)
    def _mf():
        f = input.m_file()
        if f:
            job.set((f[0]["datapath"], f[0]["name"], f[0]["size"], None))

    for btn, (_, rel, key) in BATCH_DEMOS.items():
        def _mk(btn=btn, rel=rel, key=key):
            @reactive.effect
            @reactive.event(input[btn])
            def _():
                p = SAMPLES / rel
                ui.update_select("m_ref", selected=key)
                job.set((str(p), p.name, p.stat().st_size, key))
        _mk()

    @reactive.extended_task
    async def run(path, name, size, r, splice):
        return await asyncio.to_thread(run_variant_batch, path, name, size, r, splice)

    @reactive.effect
    @reactive.event(job)
    def _go():
        if not job():
            return
        path, name, size, key = job()
        with reactive.isolate():
            r = builtin_reference(key) if key else ref()
            sp = bool(input.m_splice()) and r is not None and r.eukaryotic_nuclear and ml.available()
        if r is not None:
            run(path, name, size, r, sp)

    @render.ui
    def m_batch_out():
        st = run.status()
        if st == "initial":
            return _notice("empty", "Upload variants or try an example",
                           "Rows that are not variants of the reference are listed as not analysed, with the reason.")
        if st == "running":
            return _notice("empty", "Annotating", "Checking every variant against the reference.")
        if st == "error":
            return _notice("blocked", "The file could not be read", "Something went wrong while reading it.")
        res = run.result()
        if isinstance(res, str):
            return _notice("blocked", "This file was refused", res)
        s = res["summary"]
        log("Batch variants", f"{res['name']}: {s['n']:,} variants vs {res['ref']}",
            f"{s['annotated']:,} annotated, Ti/Tv {s['ti_tv']:.2f}" if s["ti_tv"] else f"{s['annotated']:,} annotated")
        metrics = [("variants", f"{s['n']:,}"), ("annotated", f"{s['annotated']:,}"),
                   ("not analysed", f"{s['n'] - s['annotated']:,}"),
                   ("Ti/Tv ratio", f"{s['ti_tv']:.2f}" if s["ti_tv"] else "n/a"),
                   ("likely damaging", f"{res['severe']:,}")]
        notes = [ui.p(n, class_="small-note") for n in res["notes"]]
        charts = [ui.card(ui.card_header("Effects"), ui.output_ui("m_effect_plot"))]
        if s["snv"]:
            charts.append(ui.card(ui.card_header("Single-base substitution spectrum"), ui.output_ui("m_spectrum_plot")))
        if len(s["genes"]) > 1:
            charts.append(ui.card(ui.card_header("Variants per gene"), ui.output_ui("m_gene_plot")))
        return ui.TagList(
            ui.HTML('<div class="metric-row">' + "".join(f'<div class="metric"><b>{v}</b><span>{k}</span></div>'
                                                         for k, v in metrics) + "</div>"),
            *notes,
            ui.layout_columns(*charts, col_widths=(6, 6) if len(charts) > 1 else (12,)),
            ui.card(ui.card_header(f"Variants, against {res['ref']}"), ui.output_data_frame("m_batch_table")),
            ui.download_button("m_dl", "Download annotated variants (CSV)", class_="btn-primary mt-2 w-auto"))

    @render.ui
    @_plot
    def m_effect_plot():
        res = run.result()
        req(isinstance(res, dict))
        return V.effect_bar(res["df"]["Effect"].tolist())

    @render.ui
    @_plot
    def m_spectrum_plot():
        res = run.result()
        req(isinstance(res, dict) and res["summary"]["snv"])
        return V.spectrum_bar(res["summary"]["spectrum"])

    @render.ui
    @_plot
    def m_gene_plot():
        res = run.result()
        req(isinstance(res, dict))
        return V.count_bar(res["summary"]["genes"], "Variants")

    @render.data_frame
    def m_batch_table():
        res = run.result()
        req(isinstance(res, dict))
        return render.DataGrid(res["df"], height="420px", filters=True)

    @render.download(filename="codon_bench_variants.csv")
    def m_dl():
        return run.result()["csv"]

    # ---- structural -------------------------------------------------------------------------------
    @reactive.calc
    def gene_a() -> tuple[str, str, int]:
        """(coding sequence used as gene A, its name, genetic code)."""
        r = ref()
        if r is None:
            return "", "", 1
        if r.kind == "annotated" and r.cds:
            f = max(r.cds, key=lambda x: sum(b - a for a, b in x.parts))
            return r.feature_seq(f), f.name, f.table
        return (r.seq, r.id, r.table) if len(r.seq) <= 20000 else ("", r.id, r.table)

    def _need_gene():
        dna, name, _ = gene_a()
        if not dna:
            return _notice("blocked", "No gene to rearrange", "Choose a reference with a coding sequence under 20 kb.")
        return None

    @render.ui
    def inv_out():
        if (x := _need_gene()):
            return x
        dna, name, code = gene_a()
        a, b = int(input.inv_a() or 1), int(input.inv_b() or 2)
        if not (1 <= a < b <= len(dna)):
            return _notice("blocked", "Check the positions", f"Choose 1 ≤ From < To ≤ {len(dna)}.")
        mut, e = M.inversion_effect(dna, a, b, code)
        svg = _chromosome([(name, 0, a - 1, "#1baf7a", False), (name, a - 1, b, "#1baf7a", True),
                           (name, b, len(dna), "#1baf7a", False)], len(dna), f"{name}, inverted segment hatched",
                          marks=[a - 1, b])
        return ui.TagList(ui.HTML(svg), ui.p(
            "A paracentric inversion leaves the centromere outside the flipped segment; a pericentric one includes it. "
            "No DNA is lost, but a gene that crosses a breakpoint is broken, and the flipped bases read as their "
            "reverse complement.", class_="small-note"),
            _effect_view(dna, mut, M.Mutation("inversion", a, b, dna[a - 1:b], mut[a - 1:b], f"c.{a}_{b}inv"), e,
                         is_hbb=(name == "HBB")))

    @render.ui
    def tr_out():
        if (x := _need_gene()):
            return x
        dna, name, code = gene_a()
        pb = S.parse(input.tr_b() or "")
        if pb.kind not in ("DNA", "RNA"):
            return _notice("blocked", "Partner gene not usable", pb.reason or "Gene B must be DNA or RNA.")
        b = S.to_dna(pb.seq)
        xa, xb = int(input.tr_a() or 1), int(input.tr_bb() or 1)
        if not (1 <= xa < len(dna)) or not (1 <= xb < len(b)):
            return _notice("blocked", "Check the breakpoints", f"Gene A: 1 to {len(dna) - 1}. Gene B: 1 to {len(b) - 1}.")
        t = M.translocate(dna, b, xa, xb, code)
        ref_prot_b = M._protein(b, code)
        top = max(len(dna), len(b), xa + len(b) - xb, xb + len(dna) - xa)
        svg = (_chromosome([(name, 0, xa, "#1baf7a", False), (name, xa, len(dna), "#1baf7a", False)], len(dna),
                           f"Gene A ({name}) before", marks=[xa], scale=top) +
               _chromosome([("B", 0, xb, "#4a3aa7", False), ("B", xb, len(b), "#4a3aa7", False)], len(b),
                           "Gene B before", marks=[xb], scale=top) +
               _chromosome([(name, 0, xa, "#1baf7a", False), ("B", xb, len(b), "#4a3aa7", False)], xa + len(b) - xb,
                           "der(A): start of A joined to the end of B", marks=[xa], scale=top) +
               _chromosome([("B", 0, xb, "#4a3aa7", False), (name, xa, len(dna), "#1baf7a", False)],
                           xb + len(dna) - xa, "der(B): start of B joined to the end of A", marks=[xb], scale=top))
        fusion = t.fusion_protein
        tail_b = ref_prot_b[(xb + 2) // 3:] if t.in_frame else ""
        verdict = ("In frame: the fusion protein carries B's protein after the join."
                   if t.in_frame and tail_b and fusion.endswith(tail_b) else
                   "In frame at the join." if t.in_frame else "Out of frame at the join.")
        return ui.TagList(
            ui.HTML(svg),
            ui.div(ui.h4(ui.HTML(_badge("Translocation", "severe")), " Reciprocal translocation"),
                   ui.p(M.EFFECT_INFO["Translocation"][0]), ui.p(t.partner_frame_note),
                   _facts([("Fusion gene der(A)", f"{len(t.der_a):,} bases"),
                           ("Fusion protein", f"{len(fusion.rstrip('*')):,} amino acids"),
                           ("Reading frame at the join", verdict),
                           ("Gene A's normal protein", f"{len(M._protein(dna, code).rstrip('*')):,} amino acids")]),
                   ui.h5("Fusion protein, with A's part and B's part", class_="mt-3"),
                   ui.HTML(_fusion_chips(fusion, xa)),
                   ui.p("Real example: in chronic myeloid leukaemia, t(9;22) joins BCR on chromosome 22 to ABL1 on "
                        "chromosome 9, making the BCR-ABL1 fusion kinase (the Philadelphia chromosome). The default "
                        "gene B here is a short synthetic sequence.", class_="small-note mt-2"),
                   class_="effect-card"))

    @render.ui
    def dd_out():
        if (x := _need_gene()):
            return x
        dna, name, code = gene_a()
        a, b = int(input.dd_a() or 1), int(input.dd_b() or 1)
        if not (1 <= a <= b <= len(dna)):
            return _notice("blocked", "Check the positions", f"Choose 1 ≤ From ≤ To ≤ {len(dna)}.")
        kind = input.dd_kind()
        m = M.Mutation(kind, a, b, dna[a - 1:b], dna[a - 1:b] if kind == "duplication" else "")
        m.hgvs = M.hgvs_of(m)
        mut = M.apply(dna, m)
        e = M.effect(dna, mut, m, code)
        if kind == "duplication":
            segs = [(name, 0, b, "#1baf7a", False), ("copy", a - 1, b, "#eda100", False), (name, b, len(dna), "#1baf7a", False)]
            L = len(dna) + (b - a + 1)
            title = "After the duplication: the copied segment (amber) follows the original"
        else:
            segs = [(name, 0, a - 1, "#1baf7a", False), (name, b, len(dna), "#1baf7a", False)]
            L = len(dna) - (b - a + 1)
            title = "After the deletion: the segment is gone and the ends are joined"
        top = max(len(dna), L)
        part = "to copy" if kind == "duplication" else "to delete"
        svg = (_chromosome([(name, 0, a - 1, "#1baf7a", False), (part, a - 1, b, "#eda100", False),
                            (name, b, len(dna), "#1baf7a", False)], len(dna), f"{name} before", marks=[a - 1, b],
                           scale=top) + _chromosome(segs, L, title, scale=top))
        return ui.TagList(ui.HTML(svg), _effect_view(dna, mut, m, e, is_hbb=(name == "HBB")))


# ---------------------------------------------------------------------------------------------
# Rendering helpers specific to this page
# ---------------------------------------------------------------------------------------------
def _summary_card(r: G.Reference, a: dict):
    loc = a["gene"] or ("no gene" if r.kind == "annotated" else r.id)
    items = [("Change", a["hgvs_g"] if r.kind != "coding" else a["hgvs_c"] or a["hgvs_g"]),
             ("DNA class", a["type"] + (f", {a['ti_tv'].lower()}" if a["ti_tv"] else ""))]
    if r.kind != "coding":
        items.append(("Location", f"{loc}, {a['region']}" if a["region"] else loc))
        if a["hgvs_c"]:
            items.append(("Coding change", f"{a['gene']}:{a['hgvs_c']}" if a["gene"] else a["hgvs_c"]))
    if a["hgvs_p"]:
        items.append(("Protein change", a["hgvs_p"] + (f"  ({a['protein_short']})" if a["protein_short"] else "")))
    if a.get("also"):
        items.append(("Also affects", a["also"]))
    if a.get("splice_note"):
        items.append(("ML splice prediction", a["splice_note"]))
    return ui.div(ui.h4(ui.HTML(_badge(a["effect"], a["severity"])), class_="mb-2"), _facts(items),
                  class_="effect-card")


def _noncoding_card(r: G.Reference, v: G.Variant, a: dict):
    info = M.EFFECT_INFO.get(a["effect"], ("", ""))[0]
    if not info:
        info = ("The change is in an RNA gene or regulatory region, which is not translated. Its effect depends on "
                "the RNA's structure and function." if "RNA" in a["effect"] or "region" in a["effect"].lower() else "")
    p = v.pos - 1
    lo = max(0, p - 18)
    ctx_r = r.seq[lo:p + 24]
    p0, rr, aa = v.core()
    ctx_m = r.seq[lo:p0] + aa + r.seq[p0 + len(rr):p + 24 + max(0, len(aa) - len(rr))]
    m = M.Mutation("substitution" if len(rr) == len(aa) == 1 else "delins", p0 - lo + 1, p0 - lo + max(1, len(rr)),
                   rr, aa)
    return ui.div(ui.p(info) if info else ui.TagList(),
                  ui.h5(f"DNA around {a['hgvs_g']}, reference above variant", class_="mt-2"),
                  ui.HTML(_dna_compare(ctx_r, ctx_m, m, prefix=f"{r.prefix}.", offset=lo, codons=False)),
                  class_="effect-card")


def _splice_card(a: dict):
    d = a["splice_detail"]
    rows = []
    for k in ("donor", "acceptor"):
        g, l = d[k]["gain"], d[k]["loss"]
        rows.append((f"{k.capitalize()} gain", f"{g:+.2f}" + (f" at {d[k]['gain_pos']:,}" if g > 0 else "")))
        rows.append((f"{k.capitalize()} loss", f"{l:+.2f}" + (f" at {d[k]['loss_pos']:,}" if l < 0 else "")))
    rows.append(("Strand scored", d["strand"]))
    level = "severe" if d["max_delta"] >= 0.5 else "moderate" if d["max_delta"] >= 0.2 else "neutral"
    label = ("Strong predicted splice change" if level == "severe" else "Possible splice change" if level == "moderate"
             else "No predicted splice change")
    return ui.div(
        ui.h4(ui.HTML(_badge(label, level)), class_="mb-2"),
        ui.p("The splice-site models score every GT (donor) and AG (acceptor) within 50 bases of the change, before and "
             "after it. A change of 0.2 or more is reported; 0.5 or more is strong. Scores are probabilities from "
             "models trained on mouse chromosome 19 and tested on human chromosome 21 (see Model performance).",
             class_="small-note"),
        _facts(rows), class_="effect-card")


# ---------------------------------------------------------------------------------------------
# Batch runner (runs in a worker thread)
# ---------------------------------------------------------------------------------------------
def _variants_from_table(path, kind, r: G.Reference) -> tuple[list[tuple[str, G.Variant | None, str]], list[str]]:
    """[(row id, variant or None, problem)]"""
    if kind.ext == ".xlsx":
        from openpyxl import load_workbook
        wb = load_workbook(path, read_only=True, data_only=True)
        it = wb.worksheets[0].iter_rows(values_only=True)
        header = [str(h or "") for h in next(it)]
        data = ([("" if x is None else str(x)) for x in row] for row in it)
    else:
        fh = F.open_text(path, kind.gz)
        reader = csv.reader(fh, delimiter="\t" if kind.ext == ".tsv" else ",")
        header = next(reader)
        data = reader
    low = [h.strip().lower() for h in header]

    def col(*names):
        return next((i for i, h in enumerate(low) if h in names), None)
    c_hgvs = col("hgvs", "mutation", "variant", "change", "hgvs_c", "hgvs_g", "cdna", "c_dna")
    c_pos, c_ref, c_alt = col("pos", "position", "start"), col("ref", "reference"), col("alt", "alternate", "allele")
    c_seq = col("sequence", "mutant", "mutant_sequence", "seq")
    c_id = col("id", "name", "sample", "label", "variant_id")
    out, notes = [], []
    if c_hgvs is None and None in (c_pos, c_ref, c_alt) and c_seq is None:
        raise F.Blocked("No variant columns were found. Use a column named hgvs (values like c.20A>T or g.100A>G), "
                        "columns named pos, ref and alt, or a column named sequence with mutant sequences.")
    for k, row in enumerate(data):
        if k >= 50000:
            notes.append("Only the first 50,000 rows were read.")
            break
        rid = row[c_id] if c_id is not None and c_id < len(row) and row[c_id] else f"row {k + 1}"
        try:
            if c_hgvs is not None and c_hgvs < len(row) and row[c_hgvs].strip():
                out.append((rid, G.parse_genomic_change(row[c_hgvs], r), ""))
            elif None not in (c_pos, c_ref, c_alt) and max(c_pos, c_ref, c_alt) < len(row) and row[c_pos].strip():
                out.append((rid, G.Variant(r.id, int(float(row[c_pos])), row[c_ref].strip().upper(),
                                           row[c_alt].strip().upper()), ""))
            elif c_seq is not None and c_seq < len(row) and row[c_seq].strip():
                out.append((rid, None, "__seq__" + row[c_seq]))
        except (M.BadNotation, ValueError) as e:
            out.append((rid, None, str(e)))
    return out, notes


def run_variant_batch(path, name, size, r: G.Reference, splice: bool):
    try:
        kind = F.sniff(path, name, size)
        if kind.family in ("image", "pdf"):
            return "Images and PDFs go on the 'Read a document' page."
        items: list[tuple[str, G.Variant | None, str]] = []
        notes: list[str] = []
        text = None
        if kind.family in ("variants", "sequence"):
            F.check_text_file(path, kind)
            text = G.read_file(path, name)
        if kind.family == "variants" or (text is not None and G.looks_like_vcf(text)):
            vs, problems = G.parse_vcf(text)
            if not vs and problems:
                return problems[0]
            items = [(v.id or f"{v.chrom}:{v.pos}", v, "") for v in vs]
            notes += problems[:5]
            chroms = {v.chrom for v in vs}
            if r.kind != "coding" and chroms and not any(c.split(".")[0] in (r.id.split(".")[0], "chrM", "MT", "M")
                                                         or r.id.split(".")[0] in c for c in chroms):
                notes.append(f"The VCF names its sequence {', '.join(sorted(chroms))}; positions were read against "
                             f"{r.id}. Check that this is the same reference.")
        elif kind.family == "table":
            items, n2 = _variants_from_table(path, kind, r)
            notes += n2
        elif kind.family == "sequence":
            for k, rec in enumerate(F.records(path, kind)):
                if k >= 2000:
                    notes.append("Only the first 2,000 sequences were compared.")
                    break
                items.append((rec.id, None, "__seq__" + (rec.seq or "")))
        rows = []
        for rid, v, prob in items:
            if v is not None:
                rows.append({"_id": rid, **G.annotate(r, v)})
                continue
            if prob.startswith("__seq__"):
                p = S.parse(prob[7:])
                if p.kind not in ("DNA", "RNA"):
                    rows.append({"_id": rid, "status": f"Not a DNA or RNA sequence. {p.reason}".strip()})
                    continue
                q = S.to_dna(p.seq)
                try:
                    if r.kind != "annotated" and len(r.seq) <= 20000 and len(q) <= 20000:
                        found = [G.mutation_to_variant(m, r) for m in M.detect(r.seq, q)]
                    else:
                        found = G.compare_genomes(r, q)
                except M.BadNotation as e:
                    rows.append({"_id": rid, "status": str(e)})
                    continue
                if not found:
                    rows.append({"_id": rid, "effect": "No change", "severity": "neutral",
                                 "status": "Identical to the reference"})
                for v2 in found:
                    rows.append({"_id": rid, **G.annotate(r, v2)})
                continue
            rows.append({"_id": rid, "status": prob or "Not a variant"})
        if not rows:
            return "No variants or sequences were found in this file."
        if not any(x.get("status", "").startswith("Annotated") or x.get("effect") == "No change" for x in rows):
            first = next((x["status"] for x in rows if x.get("status")), "")
            return ("Nothing in this file could be read as a variant of the reference. " + first).strip()
        if splice:
            n = 0
            for x in rows:
                if x.get("status", "").startswith("Annotated") and n < SPLICE_BATCH_LIMIT:
                    x.update({k: v for k, v in G.splice_scores(r, G.Variant(r.id, x["pos"], x["ref"], x["alt"])).items()
                              if k != "splice_detail"})
                    n += 1
            if sum(1 for x in rows if x.get("status", "").startswith("Annotated")) > SPLICE_BATCH_LIMIT:
                notes.append(f"Splice impact was scored for the first {SPLICE_BATCH_LIMIT:,} variants.")
        summary = G.summarise(rows)
        df = pd.DataFrame([{
            "ID": x.get("_id", ""), "Position": x.get("pos", ""), "REF": x.get("ref", ""), "ALT": x.get("alt", ""),
            "Type": x.get("type", ""), "Gene": x.get("gene", ""), "Region": x.get("region", ""),
            "HGVS g.": x.get("hgvs_g", ""), "HGVS c.": x.get("hgvs_c", ""), "HGVS p.": x.get("hgvs_p", ""),
            "Protein": x.get("protein_short", ""), "Effect": x.get("effect", "") or "Not analysed",
            "Splice Δ (ML)": "" if x.get("splice_max_delta") is None else x["splice_max_delta"],
            "Splice note": x.get("splice_note", ""), "Status": x.get("status", "")} for x in rows])
        if not splice:
            df = df.drop(columns=["Splice Δ (ML)", "Splice note"])
        out = tempfile.NamedTemporaryFile("w", suffix=".csv", delete=False, newline="")
        df.to_csv(out, index=False)
        out.close()
        severe = sum(1 for x in rows if x.get("severity") == "severe")
        return {"df": df, "csv": out.name, "summary": summary, "notes": notes, "severe": severe, "ref": r.id,
                "name": name}
    except F.Blocked as e:
        return str(e)
